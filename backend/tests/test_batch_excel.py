"""多根钢束 Excel 导入与批量计算测试。"""

from io import BytesIO
import unittest

from openpyxl import load_workbook

from app.services.batch_excel import calculate_project, create_template, export_project, import_project


class BatchExcelTest(unittest.TestCase):
    def test_template_can_import_calculate_and_export(self) -> None:
        """模板应可完整经过导入、计算和成果导出流程。"""
        project = import_project(create_template())
        calculation = calculate_project(project)
        exported = load_workbook(BytesIO(export_project(project, calculation)), data_only=True)

        self.assertEqual(project.tendons[0].status, "ready")
        self.assertEqual(calculation.success_count, 1)
        self.assertIn("钢束清单", exported.sheetnames)
        self.assertIn("分段明细", exported.sheetnames)
        self.assertEqual(exported["钢束清单"].cell(2, 9).value, "成功")

    def test_straight_reference_is_generated_from_other_profile(self) -> None:
        """填写“直线”时应从另一方向线形生成零 y 的两节点线。"""
        workbook = load_workbook(BytesIO(create_template()))
        workbook["钢束清单"].cell(2, 3).value = "直线"
        stream = BytesIO()
        workbook.save(stream)

        project = import_project(stream.getvalue())

        plan = project.tendons[0].plan
        self.assertEqual(plan.points[0].x, 0)
        self.assertEqual(plan.points[-1].x, 75)
        self.assertTrue(all(point.y == 0 for point in plan.points))
