"""多根钢束 Excel 导入、计算与成果导出。"""

from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.calculations import FrictionParameters, MaterialParameters, Profile, TendonElongationCalculator, TendonGeometry, TensioningCase
from app.calculations.geometry import GeometryError
from app.models.schemas import (
    BatchCalculationItem,
    BatchCalculationResponse,
    BatchDefaults,
    BatchProject,
    BatchTendon,
    CalculationResponse,
    DistributionPoint,
    NamedProfile,
    ProfileInput,
    ProfilePointInput,
    SegmentOutput,
)


STRAIGHT = "直线"
TOLERANCE_M = 0.001


class BatchExcelError(ValueError):
    """固定模板不符合要求时抛出。"""


def import_project(content: bytes) -> BatchProject:
    """读取固定工作表模板并生成可批量计算的工程对象。"""
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as error:
        raise BatchExcelError("无法读取 Excel 文件，请使用 .xlsx 格式") from error
    required = {"钢束清单", "立面线形", "平面线形", "默认参数"}
    missing = required - set(workbook.sheetnames)
    if missing:
        raise BatchExcelError(f"缺少工作表：{'、'.join(sorted(missing))}")

    defaults = _read_defaults(workbook["默认参数"])
    scale = 1.0 if defaults.unit == "m" else 0.001
    elevation_profiles = _read_profiles(workbook["立面线形"], scale)
    plan_profiles = _read_profiles(workbook["平面线形"], scale)
    tendons = _read_tendons(workbook["钢束清单"], defaults, elevation_profiles, plan_profiles)
    if not tendons:
        raise BatchExcelError("钢束清单至少需要一根钢束")
    return BatchProject(
        defaults=defaults,
        elevation_profiles=[NamedProfile(id=key, points=value.points) for key, value in elevation_profiles.items()],
        plan_profiles=[NamedProfile(id=key, points=value.points) for key, value in plan_profiles.items()],
        tendons=tendons,
    )


def calculate_project(project: BatchProject) -> BatchCalculationResponse:
    """逐根计算，单根异常不会阻断其他钢束。"""
    items: list[BatchCalculationItem] = []
    for tendon in project.tendons:
        if tendon.status != "ready" or tendon.elevation is None or tendon.plan is None:
            items.append(BatchCalculationItem(id=tendon.id, status="skipped", error=tendon.error or "导入校核未通过"))
            continue
        try:
            items.append(BatchCalculationItem(id=tendon.id, status="success", result=_calculate(tendon)))
        except (GeometryError, ValueError) as error:
            items.append(BatchCalculationItem(id=tendon.id, status="failed", error=str(error)))
    return BatchCalculationResponse(
        items=items,
        success_count=sum(item.status == "success" for item in items),
        failed_count=sum(item.status != "success" for item in items),
    )


def create_template() -> bytes:
    """生成固定列名的批量导入模板。"""
    workbook = Workbook()
    defaults = workbook.active
    defaults.title = "默认参数"
    defaults.append(["参数", "值"])
    defaults.append(["单位", "m"])
    defaults.append(["k", 0.0015])
    defaults.append(["μ", 0.25])
    defaults.append(["Ep", 195000])
    defaults.append(["左端应力", 1395])
    defaults.append(["右端应力", 1395])

    tendons = workbook.create_sheet("钢束清单")
    tendons.append(["编号", "立面编号", "平面编号", "左端应力", "右端应力", "k", "μ", "Ep"])
    tendons.append(["T1", "V1", "P1", None, None, None, None, None])
    elevation = workbook.create_sheet("立面线形")
    elevation.append(["线形编号", "x", "y", "b"])
    elevation.append(["V1", 0, 0, 0])
    elevation.append(["V1", 75, 0, 0])
    plan = workbook.create_sheet("平面线形")
    plan.append(["线形编号", "x", "y", "b"])
    plan.append(["P1", 0, 0, 0])
    plan.append(["P1", 75, 0, 0])
    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_project(project: BatchProject, calculation: BatchCalculationResponse) -> bytes:
    """导出输入、汇总结果与分段明细到新的工作簿。"""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "钢束清单"
    headers = ["编号", "立面编号", "平面编号", "左端应力 (MPa)", "右端应力 (MPa)", "k", "μ", "Ep (MPa)", "计算状态", "错误信息", "钢束总长 (m)", "累计转角 (rad)", "不动点 x (m)", "左端伸长量 (mm)", "右端伸长量 (mm)", "总伸长量 (mm)"]
    summary.append(headers)
    results = {item.id: item for item in calculation.items}
    for tendon in project.tendons:
        item = results.get(tendon.id)
        result = item.result if item else None
        summary.append([
            tendon.id, tendon.elevation_id, tendon.plan_id, tendon.left_stress, tendon.right_stress, tendon.k, tendon.mu, tendon.elastic_modulus,
            _status_text(item), item.error if item else "未计算", result.total_length if result else None, result.total_turn if result else None,
            result.balance_x if result else None, result.left_elongation_mm if result else None, result.right_elongation_mm if result else None,
            result.total_elongation_mm if result else None,
        ])

    _write_profiles(workbook.create_sheet("立面线形"), project.elevation_profiles)
    _write_profiles(workbook.create_sheet("平面线形"), project.plan_profiles)
    defaults = workbook.create_sheet("默认参数")
    defaults.append(["参数", "值"])
    for label, value in [("单位", project.defaults.unit), ("k", project.defaults.k), ("μ", project.defaults.mu), ("Ep", project.defaults.elastic_modulus), ("左端应力", project.defaults.left_stress), ("右端应力", project.defaults.right_stress)]:
        defaults.append([label, value])

    details = workbook.create_sheet("分段明细")
    details.append(["钢束编号", "伸长方向", "x 起 (m)", "x 终 (m)", "本段长度 (m)", "转角 (rad)", "平均应力 (MPa)", "本段伸长量 (mm)"])
    for item in calculation.items:
        if item.result:
            for segment in item.result.segments:
                details.append([item.id, "左端" if segment.source == "left" else "右端", segment.x_start, segment.x_end, segment.length, segment.turn, segment.average_stress, segment.elongation_mm])

    notes = workbook.create_sheet("导入与计算说明")
    notes.append(["项目", "内容"])
    notes.append(["导入单位", project.defaults.unit])
    notes.append(["计算成功", calculation.success_count])
    notes.append(["未成功", calculation.failed_count])
    notes.append(["说明", "线形编号为“直线”时，按另一方向首尾 x 自动生成。"])
    for sheet in workbook.worksheets:
        _style_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _read_defaults(sheet) -> BatchDefaults:
    values = _read_key_value_sheet(sheet)
    unit = str(values.get("单位", "m")).strip().lower()
    if unit not in {"m", "mm"}:
        raise BatchExcelError("默认参数表中的单位仅支持 m 或 mm")
    return BatchDefaults(
        unit=unit,
        k=_float(values.get("k"), 0.0015, "k"),
        mu=_float(values.get("μ"), 0.25, "μ"),
        elastic_modulus=_float(values.get("Ep"), 195000.0, "Ep"),
        left_stress=_float(values.get("左端应力"), 1395.0, "左端应力"),
        right_stress=_float(values.get("右端应力"), 1395.0, "右端应力"),
    )


def _read_profiles(sheet, scale: float) -> dict[str, ProfileInput]:
    rows = _table_rows(sheet, ["线形编号", "x", "y", "b"])
    groups: dict[str, list[ProfilePointInput]] = {}
    for row in rows:
        profile_id = _text(row["线形编号"])
        if not profile_id:
            raise BatchExcelError(f"工作表“{sheet.title}”存在空线形编号")
        groups.setdefault(profile_id, []).append(ProfilePointInput(x=_float(row["x"], None, "x") * scale, y=_float(row["y"], None, "y") * scale, value=_float(row["b"], 0.0, "b")))
    if not groups:
        raise BatchExcelError(f"工作表“{sheet.title}”未找到线形数据")
    return {profile_id: _normalize_profile(points, f"{sheet.title}：{profile_id}") for profile_id, points in groups.items()}


def _read_tendons(sheet, defaults: BatchDefaults, elevation_profiles: dict[str, ProfileInput], plan_profiles: dict[str, ProfileInput]) -> list[BatchTendon]:
    rows = _table_rows(sheet, ["编号", "立面编号", "平面编号", "左端应力", "右端应力", "k", "μ", "Ep"])
    ids: set[str] = set()
    tendons: list[BatchTendon] = []
    for row in rows:
        tendon_id = _text(row["编号"])
        if not tendon_id:
            raise BatchExcelError("钢束清单存在空编号")
        if tendon_id in ids:
            raise BatchExcelError(f"钢束编号“{tendon_id}”重复")
        ids.add(tendon_id)
        elevation_id, plan_id = _text(row["立面编号"]), _text(row["平面编号"])
        tendon = BatchTendon(
            id=tendon_id, elevation_id=elevation_id, plan_id=plan_id,
            left_stress=_float(row["左端应力"], defaults.left_stress, "左端应力"), right_stress=_float(row["右端应力"], defaults.right_stress, "右端应力"),
            k=_float(row["k"], defaults.k, "k"), mu=_float(row["μ"], defaults.mu, "μ"), elastic_modulus=_float(row["Ep"], defaults.elastic_modulus, "Ep"),
        )
        try:
            tendon.elevation, tendon.plan = _resolve_profiles(elevation_id, plan_id, elevation_profiles, plan_profiles)
            if tendon.left_stress == 0 and tendon.right_stress == 0:
                raise BatchExcelError("左右端张拉应力不能同时为 0")
        except BatchExcelError as error:
            tendon.status, tendon.error = "invalid", str(error)
        tendons.append(tendon)
    return tendons


def _resolve_profiles(elevation_id: str, plan_id: str, elevations: dict[str, ProfileInput], plans: dict[str, ProfileInput]) -> tuple[ProfileInput, ProfileInput]:
    if not elevation_id or not plan_id:
        raise BatchExcelError("立面编号和平面编号均不能为空")
    if elevation_id == STRAIGHT and plan_id == STRAIGHT:
        raise BatchExcelError("立面和平面不能同时填写“直线”")
    elevation = elevations.get(elevation_id) if elevation_id != STRAIGHT else None
    plan = plans.get(plan_id) if plan_id != STRAIGHT else None
    if elevation_id != STRAIGHT and elevation is None:
        raise BatchExcelError(f"未找到立面线形“{elevation_id}”")
    if plan_id != STRAIGHT and plan is None:
        raise BatchExcelError(f"未找到平面线形“{plan_id}”")
    if elevation is None:
        elevation = _straight_profile(plan)
    if plan is None:
        plan = _straight_profile(elevation)
    if abs(elevation.points[0].x - plan.points[0].x) > TOLERANCE_M or abs(elevation.points[-1].x - plan.points[-1].x) > TOLERANCE_M:
        raise BatchExcelError("平面与立面的首尾 x 坐标不一致")
    return elevation, plan


def _straight_profile(reference: ProfileInput) -> ProfileInput:
    return ProfileInput(mode="xyb", points=[ProfilePointInput(x=reference.points[0].x, y=0, value=0), ProfilePointInput(x=reference.points[-1].x, y=0, value=0)])


def _normalize_profile(points: list[ProfilePointInput], label: str) -> ProfileInput:
    if len(points) < 2:
        raise BatchExcelError(f"{label} 至少需要两个节点")
    if points[-1].x < points[0].x:
        points = [ProfilePointInput(x=points[old].x, y=points[old].y, value=0 if index == len(points) - 1 else -points[old - 1].value) for index, old in enumerate(range(len(points) - 1, -1, -1))]
    else:
        points[-1].value = 0
    for previous, current in zip(points, points[1:]):
        if current.x <= previous.x:
            raise BatchExcelError(f"{label} 的 x 坐标必须严格递增")
    return ProfileInput(mode="xyb", points=points)


def _calculate(tendon: BatchTendon) -> CalculationResponse:
    elevation = Profile.from_xyb_points([(point.x, point.y, point.value) for point in tendon.elevation.points])
    plan = Profile.from_xyb_points([(point.x, point.y, point.value) for point in tendon.plan.points])
    calculator = TendonElongationCalculator(TendonGeometry(elevation=elevation, plan=plan), FrictionParameters(k=tendon.k, mu=tendon.mu), MaterialParameters(elastic_modulus=tendon.elastic_modulus))
    case = TensioningCase(left_stress=tendon.left_stress or None, right_stress=tendon.right_stress or None)
    result = calculator.calculate(case)
    return CalculationResponse(total_length=result.total_length, total_turn=result.total_turn, total_elongation_mm=result.total_elongation * 1000, left_elongation_mm=result.left_elongation * 1000, right_elongation_mm=result.right_elongation * 1000, balance_x=result.balance_x, distribution=[DistributionPoint(**point) for point in calculator.sample_distribution(case, result=result, count=121)], segments=[SegmentOutput(x_start=segment.x_start, x_end=segment.x_end, length=segment.length, turn=segment.turn, average_stress=segment.average_stress, elongation_mm=segment.elongation * 1000, source=segment.source) for segment in result.segments])


def _table_rows(sheet, required: list[str]) -> list[dict[str, object]]:
    header = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if any(name not in header for name in required):
        raise BatchExcelError(f"工作表“{sheet.title}”列名必须为：{'、'.join(required)}")
    indexes = {name: header.index(name) for name in required}
    return [{name: row[index] if index < len(row) else None for name, index in indexes.items()} for row in sheet.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row)]


def _read_key_value_sheet(sheet) -> dict[str, object]:
    rows = _table_rows(sheet, ["参数", "值"])
    return {_text(row["参数"]): row["值"] for row in rows if _text(row["参数"])}


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _float(value: object, default: float | None, label: str) -> float:
    if value in (None, ""):
        if default is None:
            raise BatchExcelError(f"{label} 不能为空")
        return default
    try:
        return float(value)
    except (TypeError, ValueError) as error:
        raise BatchExcelError(f"{label} 必须为数值") from error


def _write_profiles(sheet, profiles: Iterable[NamedProfile]) -> None:
    sheet.append(["线形编号", "x", "y", "b"])
    for profile in profiles:
        for point in profile.points:
            sheet.append([profile.id, point.x, point.y, point.value])


def _status_text(item: BatchCalculationItem | None) -> str:
    return {"success": "成功", "failed": "计算失败", "skipped": "未计算"}.get(item.status if item else "", "未计算")


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="DCEEF2")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 30)
    sheet.freeze_panes = "A2"
